import argparse
import datetime as dt
import re
from pathlib import Path
from urllib.parse import urlparse

from openpyxl import load_workbook


SECRET_WORDS = (
    "парол",
    "логин",
    "доступ",
    "ключ",
    "token",
    "ftp",
    "ssh",
    "почт",
    "email",
)
PUBLIC_WORDS = (
    "домен",
    "хостинг",
    "провайдер",
    "регистратор",
    "тариф",
    "срок",
    "дата",
    "сайт",
)


def mask(value):
    text = str(value)
    return f"<скрыто, {len(text)} симв.>"


def public_summary(value):
    if isinstance(value, (dt.date, dt.datetime)):
        return value.isoformat()
    text = str(value).strip()
    urls = re.findall(r"https?://[^\s\"')]+", text)
    hosts = sorted({urlparse(url).hostname for url in urls if urlparse(url).hostname})
    if hosts:
        return "URL: " + ", ".join(hosts)
    if re.fullmatch(r"(?:[a-z0-9-]+\.)+[a-z]{2,}", text, re.IGNORECASE):
        return text
    return None


def main():
    parser = argparse.ArgumentParser()
    parser.add_argument("source", type=Path)
    args = parser.parse_args()

    workbook = load_workbook(args.source, data_only=False, read_only=False)
    print(f"Листов: {len(workbook.sheetnames)}")
    for sheet in workbook.worksheets:
        print(f"Лист: {sheet.title}; размер: {sheet.max_row}x{sheet.max_column}")
        for row in sheet.iter_rows():
            populated = [(cell.coordinate, cell.value, cell.hyperlink) for cell in row if cell.value not in (None, "")]
            if not populated:
                continue
            joined = " ".join(str(value).lower() for _, value, _ in populated)
            secret_row = any(word in joined for word in SECRET_WORDS)
            public_row = any(word in joined for word in PUBLIC_WORDS) and not secret_row
            rendered = []
            for index, (coordinate, value, hyperlink) in enumerate(populated):
                text = str(value).strip()
                summary = public_summary(value)
                if index == 0 or public_row:
                    safe = re.sub(r"(?i)(password|пароль)\s*[:=]\s*\S+", r"\1: <скрыто>", text)
                    rendered.append(f"{coordinate}={safe}")
                elif summary:
                    rendered.append(f"{coordinate}={summary}")
                else:
                    rendered.append(f"{coordinate}={mask(value)}")
                if hyperlink and hyperlink.target:
                    target = public_summary(hyperlink.target)
                    rendered.append(f"{coordinate}.link={target or '<скрыто>'}")
            print("; ".join(rendered))


if __name__ == "__main__":
    main()
