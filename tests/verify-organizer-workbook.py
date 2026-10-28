"""Read-only checks of the synthetic PHP export; never writes a workbook."""
import os
from openpyxl import load_workbook
from zipfile import ZipFile
from lxml import etree

path = os.environ['TARHAN_ORGANIZER_OUTPUT']
book = load_workbook(path)
assert book.sheetnames == ['Участники']
sheet = book.active
assert sheet.max_row == 209 and sheet.max_column == 42
assert sheet['B4'].value == 'ФИО участника'
assert sheet['B5'].value.startswith('=1+1 Иван ') and sheet['B5'].data_type == 's'
assert sheet['D5'].value == 0 and sheet['E5'].value == 0
assert sheet['D7'].value is None and sheet['E7'].value is None
assert sheet['D6'].value == 25 and sheet['E6'].value == 0.5
assert sheet['R5'].value == '+79990000000' and sheet['R5'].data_type == 's'
assert sheet['AC5'].value == '001' and sheet['AC5'].data_type == 's'
assert sheet['S5'].value.strftime('%Y-%m-%d') == '2009-01-02'
assert sheet['AG5'].value.year == 2026
assert sheet['AH4'].value == 'Год участия'
assert sheet['AH5'].value == sheet['AG5'].value.year and sheet['AH5'].data_type == 'n'
assert sheet.freeze_panes == 'A5'
assert sheet.tables['Participants'].ref == 'A4:AP209'
assert sheet.tables['Participants'].autoFilter.ref == 'A4:AP209'
assert sheet['AN4'].value == 'Сезон'
assert sheet['M5'].hyperlink.target.endswith('#essay')
assert sheet['P6'].hyperlink.target.endswith('#consent')
assert '_wpnonce' not in sheet['P6'].hyperlink.target
assert not any(cell.data_type == 'f' for row in sheet for cell in row)
with ZipFile(path) as archive:
    for name in archive.namelist():
        if name.endswith(('.xml', '.rels')):
            etree.fromstring(archive.read(name))
    xml = archive.read('xl/worksheets/sheet1.xml').decode('utf-8')
    assert '_x005F_x0041_' in xml, 'Excel escape sequence must stay literal'
print('PASS: XLSX types, Russian headers, zero/null, phones, school numbers, native dates, filters, panes, links and no formulas')
